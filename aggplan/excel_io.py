"""Excel round trip: a student template, a defensive importer, and a results export.

DESIGN
------
* The template uses EXACTLY the in-app worksheet layout (aggplan.worksheet): row 1 holds the
  column headings A..Q, January is row 2 ... December row 13. Therefore a formula typed in
  Excel (=E2/RATE, =MAX(0,C2+E2-D2-B2)) is character-for-character the formula the in-app
  grid accepts. Row 14 (totals) and everything right of column Q are ignored on import.
* The brief's constants are WORKBOOK-LEVEL DEFINED NAMES (LABOR, RATE, ...), pointing at
  numeric cells of the locked "Scenario" sheet, so Excel formulas can use the same names as
  the in-app formulas.
* A "_meta" sheet (visible but clearly labelled) carries machine-readable identification
  (schema, model version, seed, config hash ...). The importer refuses a workbook made for a
  different scenario or model so a student can never be graded against the wrong answer key.
* import_workbook() NEVER raises and never evaluates anything: the file is only read with
  openpyxl (twice: formulas and Excel's cached values), cell contents are turned into the raw
  text a student would have typed in the app, and the app's own safe evaluator
  (aggplan.formula) is used for checking. A formula the evaluator cannot handle (IF, other
  sheets, ...) falls back to Excel's cached value when one exists.
* Hostile or accidental input is bounded: compressed size <= 5 MB, uncompressed zip members
  <= 50 MB, only rows 1..60 and columns A..BH (60) are looked at.
* Sheets are protected WITHOUT a password (to stop accidental overwrites, not to hide data).
* Text written to cells that could start with '=' is stored as text (no formula injection).
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import MODEL_VERSION, engine, text, worksheet
from .formula import FUNCTIONS, NUM_PLAIN_RE, evaluate_grid
from .scenario import PRESET_LABELS, Scenario

SCHEMA = 1
MAX_BYTES = 5 * 1024 * 1024               # compressed upload cap
MAX_UNCOMPRESSED = 50 * 1024 * 1024       # zip-bomb cap (sum of member sizes)
MAX_ZIP_MEMBERS = 2000
MAX_ROWS = 60                             # importer ignores anything below this row ...
MAX_COLS = 60                             # ... or right of this column

SHEET_README, SHEET_SCENARIO, SHEET_GUIDE = "Read Me", "Scenario", "Column Guide"
SHEET_COMPARE, SHEET_META = "Compare", "_meta"
PLAN_SHEETS = {"chase": "Chase Plan", "level": "Level Plan", "hybrid": "Hybrid Plan"}

HYBRID_IN_COLS = ["Month", "Forecast Demand", "Workers", "Regular Production",
                  "Overtime Production", "Subcontract Production"]
HYBRID_CALC_COLS = ["Beginning Inventory", "Beginning Backlog", "Ending Inventory", "Backorders",
                    "Hires", "Layoffs", "Regular Labor Cost", "Hiring Cost", "Layoff Cost",
                    "Overtime Cost", "Subcontract Cost", "Holding Cost", "Backorder Cost",
                    "Total Monthly Cost"]
HYBRID_COLS = HYBRID_IN_COLS + HYBRID_CALC_COLS            # A..T

# Formula problems that mean "Excel understood it, our checker cannot" -> try cached value.
_FALLBACK_CODES = {"UNSUPPORTED", "NAME", "SYNTAX", "ARGS", "LIMIT"}
_EXCEL_ERRORS = {
    "#DIV/0!": "a division by zero or by an empty cell",
    "#REF!": "a reference to a cell that was deleted or doesn't exist",
    "#NAME?": "a name Excel doesn't know (check spelling of LABOR, RATE, ...)",
    "#VALUE!": "text where a number was expected",
    "#N/A": "a value that isn't available", "#NUM!": "an impossible number",
    "#NULL!": "an invalid cell range",
}

# --- styles ----------------------------------------------------------------------------
GIVEN_FONT = Font(color="0000FF")
GIVEN_FILL = PatternFill("solid", fgColor="FFF2CC")       # given data: blue on light yellow
INPUT_FILL = PatternFill("solid", fgColor="EAF3FB")       # cells the student fills
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E7E6E6")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

UNLOCKED = Protection(locked=False)
FMT_BOTTLES, FMT_MONEY, FMT_WORKERS = "#,##0", "$#,##0.00", "0.00#"
_WORKER_COLS = {"Workers", "Paid Labor Hours/Day", "Hires", "Layoffs"}


# =========================================================================================
# Small helpers
# =========================================================================================
def _col_fmt(name: str) -> str:
    if name.endswith("Cost"):
        return FMT_MONEY
    if name in _WORKER_COLS:
        return FMT_WORKERS
    return FMT_BOTTLES


def _put(ws, ref: str, value, *, font=None, fill=None, fmt=None, wrap=False, border=False):
    """Write one cell. Strings are stored as TEXT even if they start with '=' (never a formula)."""
    c = ws[ref]
    c.value = value
    if isinstance(value, str):
        c.data_type = "s"
    if font is not None:
        c.font = font
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = WRAP
    if border:
        c.border = BOX
    return c


def _heading_row(ws, row: int, names: Sequence[str], first_col: int = 1):
    for j, n in enumerate(names):
        _put(ws, f"{get_column_letter(first_col + j)}{row}", n, font=HEAD_FONT, fill=HEAD_FILL,
             wrap=True, border=True)


def _protect(ws):
    """Protect WITHOUT a password: stops accidental edits, students can still read/copy."""
    ws.protection.sheet = True
    ws.protection.formatColumns = False        # (False = allowed) let students resize/format
    ws.protection.formatRows = False
    ws.protection.formatCells = False


def _valid_defined_name(name: str) -> bool:
    """A usable Excel name: letters/underscore start, not a cell address (A1, XFD99, R1C1)."""
    if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", name):
        return False
    if re.match(r"^[A-Za-z]{1,3}\d+$", name) or re.match(r"^[Rr]\d*[Cc]\d*$", name):
        return False
    return True


def _sheet_title(label: str, used: Sequence[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", label).strip()[:31] or "Plan"
    t, k = base, 2
    while t in used:
        t = f"{base[:28]} {k}"
        k += 1
    return t


# =========================================================================================
# 1) TEMPLATE
# =========================================================================================
def _numeric_param_rows(scn: Scenario):
    """(label, numeric value or None, unit, number format, text-as-shown) for the Scenario
    sheet, from text.param_rows so the sheet always matches the brief."""
    p = scn.params
    keys = [("Beginning inventory", "beginning_inventory", FMT_BOTTLES),
            ("Safety stock", "safety_stock", FMT_BOTTLES),
            ("Maximum inventory", "max_inventory", FMT_BOTTLES),
            ("Output per worker-hour", "bottles_per_hour", "General"),
            ("Paid hours per full worker", "hours_per_day", "General"),
            ("Working days", "working_days", "General"),
            ("Capacity per worker", "bottles_per_worker", FMT_BOTTLES),
            ("Starting workforce", "starting_workforce", "General"),
            ("Regular labor", "regular_labor_cost", FMT_MONEY),
            ("Hiring", "hiring_cost", FMT_MONEY), ("Layoff", "layoff_cost", FMT_MONEY),
            ("Holding", "holding_cost", FMT_MONEY), ("Backorder", "backorder_cost", FMT_MONEY),
            ("Overtime limit", "overtime_pct", "0%"), ("Overtime cost", "overtime_cost", FMT_MONEY),
            ("Subcontract cost", "subcontract_cost", FMT_MONEY),
            ("Subcontract capacity", "subcontract_capacity", FMT_BOTTLES),
            ("Required on-time", "service_requirement_pct", "0.##"),
            ("Maximum December backlog", "terminal_backlog_max", FMT_BOTTLES),
            ("Shelf life", "shelf_life_months", "General"),
            ("Disposal cost", "disposal_cost", FMT_MONEY),
            ("Forecast error", "forecast_error_pct", "0%")]
    out = []
    for label, shown, unit in text.param_rows(scn):
        num, fmt = None, "General"
        for prefix, key, f in keys:
            if label.startswith(prefix):
                num, fmt = float(p[key]), f
                break
        if label.startswith("Required on-time"):
            unit = "percent (0 = no requirement)"
        out.append((label, num, unit, fmt, shown))
    return out


def _build_scenario_sheet(wb: Workbook, scn: Scenario):
    ws = wb.create_sheet(SHEET_SCENARIO)
    _put(ws, "A1", "Scenario: the given data for your plans", font=TITLE)
    _put(ws, "A2", "Blue text on yellow = given data. This sheet is locked (no password) so you "
                   "can read and copy it. Do not retype these numbers in formulas: use the "
                   "NAMES in section 3 instead (e.g. =F2*LABOR).")
    _put(ws, "A4", "1. Monthly demand forecast", font=BOLD)
    _heading_row(ws, 5, ["Month", "Forecast Demand", "Unit"])
    for i, (m, d) in enumerate(zip(scn.months, scn.demand)):
        r = 6 + i
        _put(ws, f"A{r}", m, font=GIVEN_FONT, fill=GIVEN_FILL, border=True)
        _put(ws, f"B{r}", float(d), font=GIVEN_FONT, fill=GIVEN_FILL, fmt=FMT_BOTTLES, border=True)
        _put(ws, f"C{r}", "bottles")
    last = 6 + len(scn.months)
    _put(ws, f"A{last}", "Total demand", font=BOLD, fill=TOTAL_FILL)
    _put(ws, f"B{last}", f"=SUM(B6:B{last - 1})", font=BOLD, fill=TOTAL_FILL, fmt=FMT_BOTTLES)
    _put(ws, f"C{last}", "bottles")

    r0 = last + 3
    _put(ws, f"A{r0 - 1}", "2. Parameters", font=BOLD)
    _heading_row(ws, r0, ["Parameter", "Value", "Unit", "As shown in the brief"])
    r = r0 + 1
    for label, num, unit, fmt, shown in _numeric_param_rows(scn):
        _put(ws, f"A{r}", label, border=True, wrap=True)
        if num is not None:
            _put(ws, f"B{r}", num, font=GIVEN_FONT, fill=GIVEN_FILL, fmt=fmt, border=True)
        else:
            _put(ws, f"B{r}", None, fill=GIVEN_FILL, border=True)
        _put(ws, f"C{r}", unit)
        _put(ws, f"D{r}", shown, border=True)
        r += 1

    r += 2
    _put(ws, f"A{r - 1}", "3. Names you can use in Excel formulas (identical to the in-app names)",
         font=BOLD)
    _heading_row(ws, r, ["Name", "Value", "Meaning"])
    r += 1
    meaning = {k: d for k, _, d in text.named_constant_rows(scn)}
    for name, val in scn.named_constants().items():
        if not _valid_defined_name(name):
            raise ValueError(f"'{name}' is not a valid Excel defined name.")
        _put(ws, f"A{r}", name, font=BOLD, border=True)
        _put(ws, f"B{r}", float(val), font=GIVEN_FONT, fill=GIVEN_FILL, fmt="General", border=True)
        _put(ws, f"C{r}", meaning.get(name, ""))
        wb.defined_names[name] = DefinedName(name, attr_text=f"{SHEET_SCENARIO}!$B${r}")
        r += 1
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 34
    _protect(ws)
    return ws


def _comment(text_: str) -> Comment:
    c = Comment(text_, "Aggregate Anxiety")
    c.width, c.height = 260, 110
    return c


def _plan_sheet(wb: Workbook, title: str, scn: Scenario, whole: bool):
    """Chase/Level template with EXACTLY worksheet.WS_COLS in row 1 and Jan = row 2."""
    ws = wb.create_sheet(title)
    un = worksheet.units(scn, whole)
    _heading_row(ws, 1, worksheet.WS_COLS)
    for j, name in enumerate(worksheet.WS_COLS):
        ws.cell(row=1, column=j + 1).comment = _comment(f"{name}\nUnit: {un[name]}")
        ws.column_dimensions[get_column_letter(j + 1)].width = 14 if j > 1 else 13
    ws.row_dimensions[1].height = 45
    n = len(scn.months)
    for i in range(n):
        r = worksheet.ROW_BASE + i
        _put(ws, f"A{r}", scn.months[i], font=GIVEN_FONT, fill=GIVEN_FILL, border=True)
        _put(ws, f"B{r}", float(scn.demand[i]), font=GIVEN_FONT, fill=GIVEN_FILL,
             fmt=FMT_BOTTLES, border=True)
        for j, name in enumerate(worksheet.WS_COLS[2:], start=3):
            c = ws.cell(row=r, column=j)
            c.fill, c.border, c.number_format = INPUT_FILL, BOX, _col_fmt(name)
            c.protection = UNLOCKED          # student-editable
    tr = worksheet.ROW_BASE + n                                      # row 14: totals (ignored)
    _put(ws, f"A{tr}", "Σ Totals (not imported)", font=BOLD, fill=TOTAL_FILL)
    ws[f"B{tr}"].fill = TOTAL_FILL
    for name in worksheet.SUM_COLS:
        L = worksheet.col_letter(name)
        _put(ws, f"{L}{tr}", f"=SUM({L}{worksheet.ROW_BASE}:{L}{tr - 1})", font=BOLD,
             fill=TOTAL_FILL, fmt=_col_fmt(name))
    ws.freeze_panes = "C2"
    _protect(ws)
    return ws


def _hybrid_sheet(wb: Workbook, scn: Scenario, whole: bool):
    ws = wb.create_sheet(PLAN_SHEETS["hybrid"])
    un = worksheet.units(scn, whole)
    un.update({"Overtime Production": "bottles (extra, above regular production)",
               "Subcontract Production": "bottles", "Overtime Cost": "$", "Subcontract Cost": "$"})
    _heading_row(ws, 1, HYBRID_COLS)
    for j, name in enumerate(HYBRID_COLS):
        ws.cell(row=1, column=j + 1).comment = _comment(
            f"{name}\nUnit: {un.get(name, '')}" + ("" if j < 6 else "\n(optional: not imported)"))
        ws.column_dimensions[get_column_letter(j + 1)].width = 14
    ws.row_dimensions[1].height = 45
    n = len(scn.months)
    for i in range(n):
        r = worksheet.ROW_BASE + i
        _put(ws, f"A{r}", scn.months[i], font=GIVEN_FONT, fill=GIVEN_FILL, border=True)
        _put(ws, f"B{r}", float(scn.demand[i]), font=GIVEN_FONT, fill=GIVEN_FILL,
             fmt=FMT_BOTTLES, border=True)
        for j, name in enumerate(HYBRID_COLS[2:], start=3):
            c = ws.cell(row=r, column=j)
            c.fill, c.border, c.number_format = INPUT_FILL, BOX, _col_fmt(name)
            c.protection = UNLOCKED
    tr = worksheet.ROW_BASE + n
    _put(ws, f"A{tr}", "Σ Totals (not imported)", font=BOLD, fill=TOTAL_FILL)
    for j, name in enumerate(HYBRID_COLS):
        if j >= 3 and name not in ("Beginning Inventory", "Beginning Backlog", "Ending Inventory",
                                   "Backorders", "Workers"):
            L = get_column_letter(j + 1)
            _put(ws, f"{L}{tr}", f"=SUM({L}2:{L}{tr - 1})", font=BOLD, fill=TOTAL_FILL,
                 fmt=_col_fmt(name))
    ws.freeze_panes = "C2"
    # Limits block (columns W+), read-only reminders of the plan requirements.
    p = scn.params
    _put(ws, "W1", "Limits and requirements", font=HEAD_FONT, fill=HEAD_FILL)
    _put(ws, "X1", "Value", font=HEAD_FONT, fill=HEAD_FILL)
    _put(ws, "Y1", "Unit", font=HEAD_FONT, fill=HEAD_FILL)
    limits = [
        ("Overtime limit per worker", scn.ot_limit_per_worker, "bottles/worker/month",
         f"{p['overtime_pct']:.0%} of {scn.cap_month:,.0f}; overtime only once regular capacity is full"),
        ("Subcontract capacity", float(p["subcontract_capacity"]), "bottles/month", ""),
        ("Storage limit (ending inventory)", float(p["max_inventory"]), "bottles", ""),
        ("Maximum December backlog", float(p["terminal_backlog_max"]), "bottles", ""),
        ("Minimum December inventory", 0.0, "bottles",
         f"0, or the safety stock ({p['safety_stock']:,.0f}) if your plan maintains safety stock"),
        ("Required on-time fulfillment", float(p["service_requirement_pct"]), "percent", ""),
        ("Regular production per worker", scn.cap_month, "bottles/worker/month",
         "Regular Production <= Workers x this"),
    ]
    for k, (lab, val, unit, note) in enumerate(limits, start=2):
        _put(ws, f"W{k}", lab, border=True)
        _put(ws, f"X{k}", val, font=GIVEN_FONT, fill=GIVEN_FILL, fmt="#,##0.##", border=True)
        _put(ws, f"Y{k}", unit)
        if note:
            _put(ws, f"Z{k}", note)
    ws.column_dimensions["W"].width = 34
    _protect(ws)
    return ws


def _guide_sheet(wb: Workbook, scn: Scenario, whole: bool, hybrid: bool):
    ws = wb.create_sheet(SHEET_GUIDE)
    un = worksheet.units(scn, whole)
    r = 1
    _put(ws, "A1", "Column guide (same column letters as the plan sheets)", font=TITLE)
    r = 3
    for kind in ("chase", "level"):
        _put(ws, f"A{r}", f"{PLAN_SHEETS[kind]}: columns, units and formulas", font=BOLD)
        r += 1
        _heading_row(ws, r, ["Column", "Name", "Unit", "How to calculate it"])
        r += 1
        how = {lab.split(" · ", 1)[1]: t for lab, t in text.formula_reference(scn, kind, whole)}
        for name in worksheet.WS_COLS:
            _put(ws, f"A{r}", worksheet.col_letter(name), border=True)
            _put(ws, f"B{r}", name, border=True)
            _put(ws, f"C{r}", un[name], border=True, wrap=True)
            _put(ws, f"D{r}", how.get(name, "given data (already filled in)"), border=True,
                 wrap=True)
            r += 1
        r += 1
    if hybrid:
        _put(ws, f"A{r}", f"{PLAN_SHEETS['hybrid']}: decisions to import (A to F) and optional "
                          "calculations (G to T)", font=BOLD)
        r += 1
        _heading_row(ws, r, ["Column", "Name", "Unit", "Note"])
        r += 1
        for j, name in enumerate(HYBRID_COLS):
            _put(ws, f"A{r}", get_column_letter(j + 1), border=True)
            _put(ws, f"B{r}", name, border=True)
            _put(ws, f"C{r}", un.get(name, "bottles" if "Production" in name else ""), border=True)
            _put(ws, f"D{r}", "decision: imported and checked" if 2 <= j <= 5 else
                 ("given data" if j < 2 else "optional: not imported"), border=True)
            r += 1
    for col, w in zip("ABCD", (9, 26, 42, 90)):
        ws.column_dimensions[col].width = w
    _protect(ws)


def _compare_sheet(wb: Workbook):
    ws = wb.create_sheet(SHEET_COMPARE)
    _put(ws, "A1", "Compare your plans (optional; paste your results)", font=TITLE)
    _heading_row(ws, 3, ["Metric", "Chase", "Level", "Hybrid"])
    for k, m in enumerate(["Total cost ($)", "Regular labor cost ($)", "Hiring + layoff cost ($)",
                           "Holding cost ($)", "Backorder cost ($)", "Overtime cost ($)",
                           "Subcontract cost ($)", "Highest workforce", "Lowest workforce",
                           "Highest inventory (bottles)", "December backlog (bottles)",
                           "On-time fulfillment (%)", "Meets all requirements? (yes/no)",
                           "Notes"], start=4):
        _put(ws, f"A{k}", m, border=True)
        for col in "BCD":
            c = ws[f"{col}{k}"]
            c.fill, c.border = INPUT_FILL, BOX
    ws.column_dimensions["A"].width = 36
    for col in "BCD":
        ws.column_dimensions[col].width = 18


def _readme_sheet(wb: Workbook, scn: Scenario, whole: bool, hybrid: bool, student_label: str,
                  sheet_names: List[str]):
    ws = wb.active
    ws.title = SHEET_README
    funcs = ", ".join(sorted(FUNCTIONS))
    wm = "whole workers" if whole else "partial workers (worker-equivalents)"
    lines: List[Tuple[str, Optional[Font]]] = [
        ("Aggregate Production Planning: Excel workbook", TITLE),
        ("Purpose: do your calculations in Excel, then bring the workbook back to the app. The "
         "app checks it with the SAME rules and feedback as typing into the in-app worksheet, "
         "so you learn the same method.", None),
        ("", None),
        ("STEP BY STEP", BOLD),
        ("1. Read the 'Scenario' sheet: it has your forecast and every cost/capacity value.", None),
        ("2. Work in 'Chase Plan' and 'Level Plan'" + (" (and 'Hybrid Plan')" if hybrid else "")
         + ". Fill only the light-blue cells. Row 1 holds the headings and January is row 2, "
         "exactly as in the app, so a formula like =E2/RATE or =MAX(0,C2+E2-D2-B2) is identical "
         "in Excel and in the app.", None),
        ("3. Use the NAMES LABOR, HIRE, LAYOFF, HIREHR, LAYOFFHR, HOLD, BACKORDER, OVERTIME, "
         "SUBCONTRACT, RATE, BEGIN, SAFETY, MAXINV, SHIFT, START in formulas (listed on the "
         "Scenario sheet). The 'Column Guide' sheet shows units and how each column is calculated.",
         None),
        ("4. Save the file as .xlsx (Excel Workbook). Do NOT rename sheets, columns or move "
         "rows, and keep the '_meta' sheet.", None),
        ("5. Upload the saved file in the app (Excel import). The app lists what it found and "
         "marks each cell correct or not.", None),
        ("", None),
        ("UNITS AND ASSUMPTIONS", BOLD),
        (f"Worker model of this template: {wm}. Bottles are whole units of product; money is in "
         f"dollars. Paid hours per worker per day: {scn.hpd:g}; capacity per worker: "
         f"{scn.cap_month:,.0f} bottles/month. Every column's unit is in the heading's comment "
         "and on the 'Column Guide' sheet.", None),
        (f"Formulas the checker understands: numbers, cell references (E2, $E$2), the names above, "
         f"+ - * / ^ and the functions {funcs}. If you use something else (IF, other sheets ...) "
         "the app falls back to the value Excel calculated, when the file contains it.", None),
        ("", None),
        ("WHAT THE IMPORT CHECKS", BOLD),
        ("- The file is a valid .xlsx and is not oversized; the '_meta' sheet is present.", None),
        ("- The workbook was created for YOUR current scenario (seed, configuration) and a "
         "compatible model version. Otherwise download a new template.", None),
        ("- The required sheets exist; the headings in row 1 (A to Q) are unchanged; Month and "
         "Forecast Demand are unchanged.", None),
        ("- Input cells hold numbers or formulas (no text); nothing is filled outside C2:Q13 "
         "(row 14 and columns right of Q are ignored).", None),
        ("- Blank cells produce a warning. Pasting values over formulas is fine.", None),
        ("- Cells showing Excel errors (#DIV/0!, #REF!, #NAME?) must be fixed first.", None),
        ("", None),
        ("MODEL RULES (same as in the app and the PDF report)", BOLD),
    ]
    for ln in (engine.__doc__ or "").strip().splitlines():
        if set(ln.strip()) <= {"-"}:
            continue
        lines.append((ln.rstrip(), BOLD if ln and not ln.startswith(" ") else None))
    lines += [
        ("", None),
        ("IDENTIFICATION (do not edit)", BOLD),
        (f"App model version: {MODEL_VERSION}", None),
        (f"Scenario seed: {scn.seed}   (base seed {scn.base_seed}, variant {scn.variant})", None),
        (f"Configuration hash: {scn.config_hash}", None),
        (f"Preset: {scn.preset} ({PRESET_LABELS.get(scn.preset, scn.preset)})", None),
        (f"Worker model: {'whole' if whole else 'partial'}", None),
        (f"Sheets expected by the importer: {', '.join(sheet_names)}", None),
    ]
    if student_label:
        lines.append((f"Prepared for: {student_label}", None))
    for k, (t, font) in enumerate(lines, start=1):
        _put(ws, f"A{k}", t if t else None, font=font, wrap=True)
    ws.column_dimensions["A"].width = 120
    _protect(ws)


def _meta_sheet(wb: Workbook, scn: Scenario, whole: bool, student_label: str,
                sheet_names: List[str]):
    ws = wb.create_sheet(SHEET_META)
    _put(ws, "A1", "key", font=HEAD_FONT, fill=HEAD_FILL)
    _put(ws, "B1", "value", font=HEAD_FONT, fill=HEAD_FILL)
    _put(ws, "D1", "MACHINE-READABLE IDENTIFICATION: the app reads this sheet when you upload. "
                   "Do not edit, rename or delete it.", font=BOLD)
    rows = [("schema", SCHEMA), ("model_version", MODEL_VERSION), ("seed", int(scn.seed)),
            ("base_seed", int(scn.base_seed)), ("variant", int(scn.variant)),
            ("config_hash", scn.config_hash), ("instructor_config_hash", scn.instructor_config_hash),
            ("preset", scn.preset), ("worker_model", "whole" if whole else "partial"),
            ("plan_sheets", ",".join(sheet_names)), ("student_label", student_label)]
    for k, (key, val) in enumerate(rows, start=2):
        _put(ws, f"A{k}", key)
        _put(ws, f"B{k}", val)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40
    _protect(ws)


def build_template(scn: Scenario, *, whole: bool = True, hybrid: bool = True,
                   student_label: str = "") -> bytes:
    """The student workbook (.xlsx bytes) for this scenario."""
    wb = Workbook()
    sheet_names = [PLAN_SHEETS["chase"], PLAN_SHEETS["level"]] + (
        [PLAN_SHEETS["hybrid"]] if hybrid else [])
    _readme_sheet(wb, scn, whole, hybrid, student_label, sheet_names)
    _build_scenario_sheet(wb, scn)
    _plan_sheet(wb, PLAN_SHEETS["chase"], scn, whole)
    _plan_sheet(wb, PLAN_SHEETS["level"], scn, whole)
    if hybrid:
        _hybrid_sheet(wb, scn, whole)
    _guide_sheet(wb, scn, whole, hybrid)
    _compare_sheet(wb)
    _meta_sheet(wb, scn, whole, student_label, sheet_names)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =========================================================================================
# 2) IMPORT
# =========================================================================================
@dataclass
class ImportMessage:
    level: str            # 'error' | 'warning' | 'info'
    sheet: str
    cell: str
    text: str

    def __str__(self):
        where = f"{self.sheet}!{self.cell}" if self.cell else self.sheet
        return f"[{self.level}] {where + ': ' if where else ''}{self.text}"


@dataclass
class ImportResult:
    ok: bool = True
    messages: List[ImportMessage] = field(default_factory=list)
    sheets: Dict[str, pd.DataFrame] = field(default_factory=dict)
    hybrid: Optional[dict] = None
    meta: Dict[str, str] = field(default_factory=dict)

    @property
    def errors(self) -> List[ImportMessage]:
        return [m for m in self.messages if m.level == "error"]

    @property
    def warnings(self) -> List[ImportMessage]:
        return [m for m in self.messages if m.level == "warning"]


class _Msgs:
    def __init__(self):
        self.items: List[ImportMessage] = []

    def add(self, level, sheet, cell, text_):
        self.items.append(ImportMessage(level, sheet, cell, text_))

    def n_errors(self):
        return sum(m.level == "error" for m in self.items)


def _fail(msgs: _Msgs, sheet: str, text_: str, cell: str = "") -> ImportResult:
    msgs.add("error", sheet, cell, text_)
    return ImportResult(ok=False, messages=msgs.items)


def _check_container(data: bytes, msgs: _Msgs) -> bool:
    """Size and zip-bomb checks done BEFORE openpyxl parses anything."""
    if not isinstance(data, (bytes, bytearray)) or len(data) == 0:
        msgs.add("error", "", "", "The uploaded file is empty. Upload the .xlsx workbook you saved.")
        return False
    if len(data) > MAX_BYTES:
        msgs.add("error", "", "", f"The file is {len(data) / 1048576:.1f} MB; the limit is "
                 f"{MAX_BYTES // 1048576} MB. Upload the template you downloaded, not a different file.")
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as z:
            infos = z.infolist()
            if len(infos) > MAX_ZIP_MEMBERS:
                msgs.add("error", "", "", "The file contains too many parts to be a normal "
                         "Excel workbook.")
                return False
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                msgs.add("error", "", "", "The file expands to far more data than a normal Excel "
                         "workbook (over 50 MB uncompressed) and was rejected for safety.")
                return False
    except zipfile.BadZipFile:
        msgs.add("error", "", "", "This is not a valid .xlsx file (it is corrupt, or it is an older "
                 ".xls/.csv file). In Excel use File > Save As > 'Excel Workbook (*.xlsx)' and "
                 "upload that.")
        return False
    return True


def _grid(ws, nrows=MAX_ROWS, ncols=MAX_COLS) -> List[list]:
    """The top-left nrows x ncols block as a list of row lists (padded with None)."""
    rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=nrows, min_col=1, max_col=ncols,
                                          values_only=True)]
    rows = [r + [None] * (ncols - len(r)) for r in rows]
    return rows + [[None] * ncols for _ in range(nrows - len(rows))]


def _norm_id(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


def _cell_ref(i: int, j: int) -> str:
    """Excel address for 0-based grid row i (January = 0) and 0-based column j."""
    return f"{get_column_letter(j + 1)}{i + worksheet.ROW_BASE}"


def _short(s: str, n: int = 60) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


def _fmt_number(v) -> str:
    f = float(v)
    return str(int(f)) if f.is_integer() and abs(f) < 1e15 else repr(f)


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _read_cells(gf, gv, sheet, ncols, strict_upto, msgs: _Msgs):
    """Turn rows 2..13 of a plan sheet into raw student text.

    Returns (raw[i][j] strings, cached {(i, j): float}). Columns with index < strict_upto
    (and >= 2) are inputs and must hold a number or formula; later columns (hybrid
    calculations) are read leniently. Problems are reported in `msgs`."""
    n = 12                                               # January..December
    raw = [[""] * ncols for _ in range(n)]
    cached: Dict[Tuple[int, int], float] = {}
    for i in range(n):
        for j in range(2, ncols):
            v = gf[worksheet.ROW_BASE - 1 + i][j]
            cv = gv[worksheet.ROW_BASE - 1 + i][j]
            ref = _cell_ref(i, j)
            strict = j < strict_upto
            if hasattr(v, "text") and not isinstance(v, str):       # array formula object
                v = "=" + str(v.text).lstrip("=")
            if v is None:
                continue
            if isinstance(v, str):
                s = v.strip()
                if s == "":
                    continue
                if s in _EXCEL_ERRORS:
                    if strict:
                        msgs.add("error", sheet, ref, f"{ref} contains the Excel error {s}. "
                                 "Fix or clear this cell and save the workbook again.")
                    continue
                if s.startswith("="):
                    raw[i][j] = s
                    if isinstance(cv, str) and cv.strip() in _EXCEL_ERRORS:
                        if strict:
                            why = _EXCEL_ERRORS[cv.strip()]
                            msgs.add("error", sheet, ref, f"{ref} ({_short(s)}) shows the Excel "
                                     f"error {cv.strip()}: {why}. Fix the formula in Excel, then save.")
                    elif _is_number(cv):
                        cached[(i, j)] = float(cv)
                elif NUM_PLAIN_RE.match(s):
                    raw[i][j] = s                           # e.g. "10,500" typed as text
                elif strict:
                    msgs.add("error", sheet, ref, f"{ref} contains the text '{_short(s, 40)}'. "
                             "Input cells must hold a number or a formula starting with '='. "
                             "Delete the text or replace it with a number.")
                continue
            if isinstance(v, bool):
                if strict:
                    msgs.add("error", sheet, ref, f"{ref} contains TRUE/FALSE; enter a number "
                             "or a formula.")
                continue
            if _is_number(v):
                raw[i][j] = _fmt_number(v)
                continue
            if strict:
                msgs.add("error", sheet, ref, f"{ref} holds something that is not a number or a "
                         "formula (a date, time or other object). Enter a plain number.")
    return raw, cached


def _check_fixed(gf, gv, sheet, scn: Scenario, ncols_head: int, head_names: Sequence[str],
                 msgs: _Msgs) -> None:
    """Row-1 headings A..; Month and Forecast Demand cells must be untouched."""
    bad = []
    for j, exp in enumerate(head_names[:ncols_head]):
        got = gf[0][j]
        got_s = "" if got is None else str(got).strip()
        if got_s != exp:
            bad.append((j, exp, got_s))
    for j, exp, got in bad[:6]:
        L = get_column_letter(j + 1)
        msgs.add("error", sheet, f"{L}1", f"{L}1 should say '{exp}' but says '{got}'. Don't rename, "
                 "move, insert or delete columns or rows. Download a fresh template and copy your "
                 "work into it if needed.")
    if len(bad) > 6:
        msgs.add("error", sheet, "", f"{len(bad) - 6} more heading cells in row 1 differ from the "
                 "template (a column or row may have been inserted or deleted).")
    for i, month in enumerate(scn.months):
        r = worksheet.ROW_BASE + i
        a = gf[r - 1][0]
        a_s = "" if a is None else str(a).strip()
        if a_s != month:
            msgs.add("error", sheet, f"A{r}", f"A{r} should be '{month}' but is '{a_s}'. The month "
                     "column is given data: don't change or reorder it.")
        b = gf[r - 1][1]
        if isinstance(b, str) and b.startswith("="):
            b = gv[r - 1][1]
        exp = float(scn.demand[i])
        if not _is_number(b) or abs(float(b) - exp) > 1e-6:
            got = f"{float(b):,.0f}" if _is_number(b) else (f"'{b}'" if b not in (None, "") else "blank")
            msgs.add("error", sheet, f"B{r}", f"B{r} (Forecast Demand for {month}) should be "
                     f"{exp:,.0f} but is {got}. The forecast is given data: don't change it.")


def _check_outside(gf, sheet, ncols, msgs: _Msgs) -> None:
    """Anything in rows 15..60 (within the sheet's column range) is outside the input area."""
    first_free = worksheet.ROW_BASE + 12 + 1            # row 15 (row 14 = totals, ignored)
    hits = []
    for r in range(first_free, MAX_ROWS + 1):
        for j in range(ncols):
            v = gf[r - 1][j]
            if v is not None and not (isinstance(v, str) and v.strip() == ""):
                hits.append(f"{get_column_letter(j + 1)}{r}")
    if hits:
        shown = ", ".join(hits[:8]) + (f" and {len(hits) - 8} more" if len(hits) > 8 else "")
        msgs.add("error", sheet, hits[0], f"Cells filled outside the input area C2:Q13 "
                 f"({shown}). Only rows 2 to 13 are read; move or delete these cells (notes "
                 "can go on another sheet).")


def _evaluate_with_fallback(df: pd.DataFrame, scn: Scenario, cached, sheet: str, msgs: _Msgs,
                            report_cols: int, *, unusable_is_error: bool):
    """Evaluate the grid; where the in-app evaluator cannot handle a formula but Excel stored
    its result, replace the cell by that number. Returns the final GridResult.
    df has a RangeIndex 0..11 (callers re-index afterwards)."""
    const = scn.named_constants()
    cols = list(df.columns)
    res = evaluate_grid(df, const, row_base=worksheet.ROW_BASE)
    replaced = set()
    for _ in range(8):                           # dependents may clear after a replacement
        changed = False
        for (c, i), e in list(res.errors.items()):
            j = cols.index(c)
            if e.code in _FALLBACK_CODES and (i, j) in cached and (i, j) not in replaced:
                f = df.at[i, c]
                df.at[i, c] = _fmt_number(cached[(i, j)])
                replaced.add((i, j))
                changed = True
                if j < report_cols:
                    ref = _cell_ref(i, j)
                    msgs.add("info", sheet, ref, f"{ref} contains '{_short(str(f))}' which the "
                             f"checker can't evaluate ({e.message.rstrip('.')}). Used Excel's "
                             f"calculated value for {ref} ({_fmt_number(cached[(i, j)])}) instead.")
        if not changed:
            break
        res = evaluate_grid(df, const, row_base=worksheet.ROW_BASE)
    for (c, i), e in res.errors.items():
        j = cols.index(c)
        raw = res.raw.get((c, i), "")
        if e.code in _FALLBACK_CODES and raw.startswith("=") and j < report_cols:
            ref = _cell_ref(i, j)
            body = (f"{ref} contains '{_short(raw)}' which the checker can't evaluate "
                    f"({e.message.rstrip('.')}) and the workbook has no saved calculated value. ")
            if unusable_is_error:
                msgs.add("error", sheet, ref, body + "Replace it with a simpler formula (MAX/MIN) "
                         "or paste its value, then save.")
            else:
                msgs.add("warning", sheet, ref, body + "Open the file in Excel (so it calculates), "
                         "save it again, or replace the formula with MAX/MIN or paste its value.")
    return res


def _frame(scn: Scenario, columns: Sequence[str], raw) -> pd.DataFrame:
    data = {"Month": list(scn.months), "Forecast Demand": [float(d) for d in scn.demand]}
    for j, c in enumerate(columns):
        if j >= 2:
            data[c] = [raw[i][j] for i in range(len(scn.months))]
    return pd.DataFrame(data)[list(columns)]


def _read_plan(kind: str, sheet: str, gf, gv, scn: Scenario, msgs: _Msgs,
               result: ImportResult) -> None:
    before = msgs.n_errors()
    cols = worksheet.WS_COLS
    _check_fixed(gf, gv, sheet, scn, len(cols), cols, msgs)
    _check_outside(gf, sheet, len(cols), msgs)
    raw, cached = _read_cells(gf, gv, sheet, len(cols), len(cols), msgs)
    df = _frame(scn, cols, raw)
    _evaluate_with_fallback(df, scn, cached, sheet, msgs, len(cols), unusable_is_error=False)
    blanks = sum(df.iat[i, j] == "" for i in range(len(df)) for j in range(2, len(cols)))
    total = len(df) * (len(cols) - 2)
    label = {"chase": "chase", "level": "level"}[kind]
    if blanks == total:
        msgs.add("warning", sheet, "", f"The {label} plan sheet is completely blank. Fill C2:Q13 "
                 "(or upload only the plans you have finished).")
    elif blanks:
        msgs.add("warning", sheet, "", f"{blanks} of {total} input cells are still blank; "
                 "blank cells count as not done (and as 0 in formulas that use them).")
    if msgs.n_errors() == before:
        df.index = range(worksheet.ROW_BASE, worksheet.ROW_BASE + len(scn.months))
        result.sheets[kind] = df


def _read_hybrid(sheet: str, gf, gv, scn: Scenario, msgs: _Msgs, result: ImportResult) -> None:
    before = msgs.n_errors()
    _check_fixed(gf, gv, sheet, scn, len(HYBRID_IN_COLS), HYBRID_IN_COLS, msgs)
    _check_outside(gf, sheet, len(HYBRID_COLS), msgs)
    raw, cached = _read_cells(gf, gv, sheet, len(HYBRID_COLS), len(HYBRID_IN_COLS), msgs)
    # Lenient calculation columns: ignore anything that is not a number/formula.
    for i in range(len(raw)):
        for j in range(len(HYBRID_IN_COLS), len(HYBRID_COLS)):
            s = raw[i][j]
            if s and not (s.startswith("=") or NUM_PLAIN_RE.match(s)):
                raw[i][j] = ""
    df = _frame(scn, HYBRID_COLS, raw)
    res = _evaluate_with_fallback(df, scn, cached, sheet, msgs, len(HYBRID_IN_COLS),
                                  unusable_is_error=True)
    keys = {"Workers": "workers", "Regular Production": "regular",
            "Overtime Production": "overtime", "Subcontract Production": "subcontract"}
    out = {k: [] for k in keys.values()}
    for name, key in keys.items():
        j = HYBRID_COLS.index(name)
        for i in range(len(scn.months)):
            r = res.raw.get((name, i), "")
            if r == "":
                out[key].append(None)
                continue
            err = res.errors.get((name, i))
            v = res.values[name][i]
            ref = _cell_ref(i, j)
            if err is not None or v is None:
                if not any(m.cell == ref and m.level == "error" for m in msgs.items):
                    msgs.add("error", sheet, ref, f"{ref} can't be calculated: "
                             f"{err.message if err else 'no value'} Fix the cell and save again.")
                out[key].append(None)
            else:
                if v < 0:
                    msgs.add("warning", sheet, ref, f"{ref} is negative ({v:,.2f}); "
                             "decisions normally can't be below 0.")
                out[key].append(float(v))
    if all(v is None for k in ("workers", "regular") for v in out[k]):
        msgs.add("warning", sheet, "", "The hybrid plan has no Workers or Regular Production "
                 "entries yet.")
    if msgs.n_errors() == before:
        result.hybrid = out


def import_workbook(data: bytes, scn: Scenario, *, whole: bool = True,
                    kinds=("chase", "level", "hybrid")) -> ImportResult:
    """Read a student workbook. Never raises; see module docstring for the rules."""
    msgs = _Msgs()
    try:
        return _import(data, scn, whole, tuple(kinds), msgs)
    except Exception as exc:                                 # defensive: never crash the app
        return _fail(msgs, "", f"The workbook could not be read ({type(exc).__name__}). Make sure "
                     "it is an unmodified .xlsx saved from the template, or download a new template.")


def _import(data, scn: Scenario, whole: bool, kinds, msgs: _Msgs) -> ImportResult:
    if not _check_container(data, msgs):
        return ImportResult(ok=False, messages=msgs.items)
    wf = wv = None
    try:
        wf = load_workbook(io.BytesIO(bytes(data)), read_only=True, data_only=False)
        wv = load_workbook(io.BytesIO(bytes(data)), read_only=True, data_only=True)
        names = list(wf.sheetnames)
        if SHEET_META not in names:
            return _fail(msgs, SHEET_META, "The '_meta' sheet is missing, so the app can't tell "
                         "which scenario this workbook belongs to. Don't delete it; or download a "
                         "new template and copy your work into it.")
        meta: Dict[str, str] = {}
        for row in _grid(wf[SHEET_META], 60, 2):
            if row[0] not in (None, ""):
                meta[str(row[0]).strip()] = _norm_id(row[1])
        result = ImportResult(meta=meta, messages=msgs.items)
        if meta.get("schema") != str(SCHEMA):
            return _fail(msgs, SHEET_META, f"This workbook uses template format "
                         f"'{meta.get('schema', 'unknown')}' but the app expects format {SCHEMA}. "
                         "Download a new template.")
        mv, cur = meta.get("model_version", ""), str(MODEL_VERSION)
        if mv.split(".")[0] != cur.split(".")[0]:
            return _fail(msgs, SHEET_META, f"This workbook was made with model version "
                         f"'{mv or 'unknown'}' but the app is on version {cur}; the calculation "
                         "rules differ. Download a new template.")
        if (meta.get("seed") != _norm_id(scn.seed) or meta.get("variant") != _norm_id(scn.variant)
                or meta.get("config_hash") != scn.config_hash):
            return _fail(msgs, SHEET_META,
                         f"This workbook was created for scenario seed {meta.get('seed', '?')} "
                         f"(config {meta.get('config_hash', '?')}) but your current scenario is "
                         f"seed {scn.seed} (config {scn.config_hash}). Download a new template for "
                         "your current scenario.")
        want = "whole" if whole else "partial"
        if meta.get("worker_model") not in (None, "", want):
            return _fail(msgs, SHEET_META, f"This template was built for "
                         f"{meta.get('worker_model')} workers but the app is set to {want} "
                         "workers. Download a template for the worker model you are using.")
        listed = [s for s in meta.get("plan_sheets", "").split(",") if s]
        for kind in kinds:
            sheet = PLAN_SHEETS.get(kind)
            if sheet is None:
                continue
            if sheet not in names:
                if sheet in listed or kind != "hybrid":
                    msgs.add("error", sheet, "", f"The required sheet '{sheet}' is missing (was it "
                             "renamed or deleted?). Restore it or download a new template.")
                else:
                    msgs.add("info", sheet, "", "This template has no Hybrid Plan sheet; "
                             "the hybrid plan was skipped.")
                continue
            gf, gv = _grid(wf[sheet]), _grid(wv[sheet])
            if kind == "hybrid":
                _read_hybrid(sheet, gf, gv, scn, msgs, result)
            else:
                _read_plan(kind, sheet, gf, gv, scn, msgs, result)
        result.ok = msgs.n_errors() == 0
        return result
    except zipfile.BadZipFile:
        return _fail(msgs, "", "This is not a valid .xlsx file. Save it from Excel as "
                     "'Excel Workbook (*.xlsx)' and upload again.")
    except Exception:
        return _fail(msgs, "", "The file could not be opened as an Excel workbook (it may be "
                     "corrupt or saved in another format). Save it as .xlsx and try again.")
    finally:
        for w in (wf, wv):
            try:
                if w is not None:
                    w.close()
            except Exception:
                pass


# =========================================================================================
# 3) EXPORT OF COMPLETED PLANS
# =========================================================================================
_SUM_EXPORT = ["Forecast Demand", "Regular Capacity", "Regular Production", "Idle Capacity",
               "Overtime Production", "Subcontract Production", "Total Production", "Hires",
               "Layoffs", "On-Time Shipments", "Backlog Cleared", "New Shortage", "Disposed",
               "Regular Labor Cost", "Hiring Cost", "Layoff Cost", "Overtime Cost",
               "Subcontract Cost", "Holding Cost", "Backorder Cost", "Disposal Cost",
               "Total Monthly Cost"]
_EXPORT_UNITS = {
    "Month": "calendar month", "Forecast Demand": "bottles", "Beginning Inventory": "bottles",
    "Beginning Backlog": "bottles", "Workers": "workers (worker-equivalents in partial mode)",
    "Paid Labor Hours/Day": "paid hours/day, all workers", "Regular Capacity": "bottles",
    "Regular Production": "bottles", "Idle Capacity": "bottles", "Overtime Production": "bottles",
    "Subcontract Production": "bottles", "Total Production": "bottles", "Hires": "workers (or hours/day)",
    "Layoffs": "workers (or hours/day)", "On-Time Shipments": "bottles", "Backlog Cleared": "bottles",
    "New Shortage": "bottles", "Ending Inventory": "bottles", "Backorders": "bottles",
    "Disposed": "bottles", "Storage Heuristic": "text note",
}


def _metric_fmt(key: str) -> str:
    k = key.lower()
    if "cost" in k:
        return FMT_MONEY
    if "fulfill" in k:
        return '0.0"%"'
    if "workforce" in k and "number" not in k or k in ("total hires", "total layoffs"):
        return FMT_WORKERS
    return FMT_BOTTLES


def _status_text(checks: Sequence[dict]) -> str:
    failed = [c["label"] for c in checks if not c["ok"]]
    return "Meets all requirements" if not failed else "Does NOT meet: " + "; ".join(failed)


def export_workbook(scn: Scenario, plans: Dict[str, dict], *, whole: bool = True,
                    student_label: str = "", status: str = "") -> bytes:
    """Completed plans -> .xlsx bytes: Summary, Assumptions and one sheet per plan.
    Stored values are the engine's unrounded numbers (formats only change the display)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    labels = list(plans)
    used = ["Summary", "Assumptions"]
    sheet_for = {}
    for lab in labels:
        sheet_for[lab] = _sheet_title(f"{lab} Plan", used)
        used.append(sheet_for[lab])

    # ---- Summary ------------------------------------------------------------------------
    _put(ws, "A1", "Plan summary: Juicetification: Aggregate Anxiety", font=TITLE)
    info = [f"Student: {student_label}" if student_label else "", f"Status: {status}" if status else "",
            f"Scenario seed {scn.seed}, configuration {scn.config_hash}, model version "
            f"{MODEL_VERSION}"]
    r = 2
    for t in info:
        if t:
            _put(ws, f"A{r}", t)
            r += 1
    r += 1
    _heading_row(ws, r, ["Metric"] + labels + ["Definition"])
    r += 1
    _put(ws, f"A{r}", "Feasibility", font=BOLD, border=True)
    for k, lab in enumerate(labels):
        _put(ws, f"{get_column_letter(2 + k)}{r}", _status_text(plans[lab].get("checks", [])),
             wrap=True, border=True)
    _put(ws, f"{get_column_letter(2 + len(labels))}{r}", "Plain-text result of every requirement "
         "check (see the check list below).", wrap=True)
    r += 1
    keys: List[str] = []
    for lab in labels:
        for k in plans[lab].get("summary", {}):
            if k not in keys:
                keys.append(k)
    for key in keys:
        _put(ws, f"A{r}", key, border=True)
        for k, lab in enumerate(labels):
            v = plans[lab].get("summary", {}).get(key)
            _put(ws, f"{get_column_letter(2 + k)}{r}", v, fmt=_metric_fmt(key), border=True)
        _put(ws, f"{get_column_letter(2 + len(labels))}{r}",
             engine.METRIC_DEFINITIONS.get(key, ""), wrap=True)
        r += 1
    r += 1
    _put(ws, f"A{r}", "Requirement checks", font=BOLD)
    r += 1
    _heading_row(ws, r, ["Plan", "Requirement", "Result", "Detail"])
    r += 1
    for lab in labels:
        for c in plans[lab].get("checks", []):
            _put(ws, f"A{r}", lab)
            _put(ws, f"B{r}", c["label"], wrap=True)
            _put(ws, f"C{r}", "OK" if c["ok"] else "NOT MET")
            _put(ws, f"D{r}", c.get("detail", ""), wrap=True)
            r += 1
    r += 1
    _put(ws, f"A{r}", "Settings per plan", font=BOLD)
    r += 1
    for lab in labels:
        for k, v in plans[lab].get("settings", {}).items():
            _put(ws, f"A{r}", lab)
            _put(ws, f"B{r}", str(k))
            _put(ws, f"C{r}", str(v))
            r += 1
    ws.column_dimensions["A"].width = 34
    for k in range(len(labels)):
        ws.column_dimensions[get_column_letter(2 + k)].width = 26
    ws.column_dimensions[get_column_letter(2 + len(labels))].width = 80
    ws.freeze_panes = "B1"

    # ---- Assumptions --------------------------------------------------------------------
    wa = wb.create_sheet("Assumptions")
    _put(wa, "A1", "Assumptions and scenario identification", font=TITLE)
    _heading_row(wa, 3, ["Parameter", "Value", "Unit"])
    r = 4
    for label, value, unit in text.param_rows(scn):
        _put(wa, f"A{r}", label, wrap=True)
        _put(wa, f"B{r}", value)
        _put(wa, f"C{r}", unit)
        r += 1
    r += 1
    _put(wa, f"A{r}", "Scenario identification", font=BOLD)
    r += 1
    ident = [("Scenario seed", scn.seed), ("Base seed", scn.base_seed), ("Variant", scn.variant),
             ("Configuration hash", scn.config_hash),
             ("Instructor configuration hash", scn.instructor_config_hash),
             ("Model version", MODEL_VERSION),
             ("Preset", f"{scn.preset} ({PRESET_LABELS.get(scn.preset, scn.preset)})"),
             ("Worker model", "whole workers" if whole else "partial workers (worker-equivalents)")]
    for k, v in ident:
        _put(wa, f"A{r}", k)
        _put(wa, f"B{r}", v)
        r += 1
    if scn.notices:
        r += 1
        _put(wa, f"A{r}", "Notices", font=BOLD)
        r += 1
        for nt in scn.notices:
            _put(wa, f"A{r}", nt, wrap=True)
            r += 1
    wa.column_dimensions["A"].width = 56
    wa.column_dimensions["B"].width = 30
    wa.column_dimensions["C"].width = 28

    # ---- one sheet per plan -------------------------------------------------------------
    for lab in labels:
        df = plans[lab]["df"]
        wp = wb.create_sheet(sheet_for[lab])
        _heading_row(wp, 1, engine.PLAN_COLUMNS)
        wp.row_dimensions[1].height = 45
        n = len(df)
        for i in range(n):
            for j, name in enumerate(engine.PLAN_COLUMNS):
                v = df[name].iloc[i]
                if name in ("Month", "Storage Heuristic"):
                    _put(wp, f"{get_column_letter(j + 1)}{i + 2}", "" if v is None else str(v))
                else:
                    _put(wp, f"{get_column_letter(j + 1)}{i + 2}", float(v), fmt=_col_fmt(name))
        tr = n + 2
        _put(wp, f"A{tr}", "Σ Total", font=BOLD, fill=TOTAL_FILL)
        for name in _SUM_EXPORT:
            L = get_column_letter(engine.PLAN_COLUMNS.index(name) + 1)
            _put(wp, f"{L}{tr}", f"=SUM({L}2:{L}{n + 1})", font=BOLD, fill=TOTAL_FILL,
                 fmt=_col_fmt(name))
        _put(wp, f"A{tr + 2}", "Units", font=BOLD)
        for j, name in enumerate(engine.PLAN_COLUMNS):
            u = "$" if name.endswith("Cost") else _EXPORT_UNITS.get(name, "")
            _put(wp, f"{get_column_letter(j + 1)}{tr + 3}", u)
        _put(wp, f"A{tr + 5}", "Values are stored at full precision; number formats only change "
                               "the display. Σ cells are SUM formulas (stocks such as inventory "
                               "and workers are not summed).")
        wp.freeze_panes = "B2"
        for j in range(len(engine.PLAN_COLUMNS)):
            wp.column_dimensions[get_column_letter(j + 1)].width = 14
        wp.column_dimensions[get_column_letter(len(engine.PLAN_COLUMNS))].width = 34
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
