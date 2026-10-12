"""The 'Excel Skills' teaching sheet, the plan-sheet colour rules and the plan-sheet chart.

The whole point of the Excel Skills sheet is that it teaches but is never graded, so most of
these tests are about what it does NOT do to the import.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import load_workbook

from aggplan import excel_io, worksheet
from aggplan.excel_io import SHEET_SKILLS, build_template, import_workbook
from tests.helpers import SEASONAL, scn

S = scn(SEASONAL)
PLANS = ("Chase Plan", "Level Plan")


# --------------------------------------------------------------------------- helpers
def _wb(**kw):
    return load_workbook(io.BytesIO(build_template(S, **kw)))


def _find_row(ws, col: str, needle: str) -> int:
    """First row whose cell in `col` equals `needle` (the sheet is laid out with a cursor)."""
    for r in range(1, ws.max_row + 1):
        if ws[f"{col}{r}"].value == needle:
            return r
    raise AssertionError(f"{needle!r} not found in column {col} of {ws.title}")


def _exercise_rows(ws):
    head = _find_row(ws, "D", "Month")
    return head, head + 1, head + len(excel_io.SKILLS_QUANTITIES)


def _text(ws, max_col=12) -> str:
    return "\n".join(str(c.value) for row in ws.iter_rows(max_col=max_col) for c in row
                     if c.value is not None)


# --------------------------------------------------------------------------- the sheet itself
def test_skills_sheet_exists_right_after_read_me():
    wb = _wb()
    assert wb.sheetnames[0] == "Read Me"
    assert wb.sheetnames[1] == SHEET_SKILLS
    assert SHEET_SKILLS not in excel_io.PLAN_SHEETS.values()


def test_skills_sheet_is_not_a_plan_sheet_in_meta():
    wb = _wb()
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(min_row=2, max_col=2)}
    assert SHEET_SKILLS not in str(meta["plan_sheets"])
    assert str(meta["plan_sheets"]).split(",") == ["Chase Plan", "Level Plan", "Hybrid Plan"]


def test_skills_sheet_teaches_every_required_block():
    body = _text(_wb()[SHEET_SKILLS])
    for marker in ("HOW A FORMULA WORKS", "FILL / COPY DOWN", "RELATIVE vs ABSOLUTE",
                   "THE FUNCTIONS THE APP ACCEPTS", "WHEN A CELL SHOWS AN ERROR", "NAMED CELLS"):
        assert marker in body, marker
    # (a) formula basics, (b) copy down
    assert "starts with an equals sign" in body and "formula bar" in body
    assert "fill handle" in body and "Ctrl+D" in body
    # (c) the four reference forms, each explained on its own row of the $ table
    ws = _wb()[SHEET_SKILLS]
    col_a = [ws[f"A{r}"].value for r in range(1, ws.max_row + 1)]
    for ref in ("A2", "$A$2", "A$2", "$A2"):
        assert ref in col_a, ref
    # (d) every function the checker accepts, with an example
    for name, _desc, example in excel_io.SKILL_FUNCTIONS:
        assert name in body and example in body
    assert "AVERAGE" in body and "COUNT" in body and "IF" in body
    assert "=SUM(E2:E13)" in body and "=SUM(L2:P2)" in body and "=SUM(L2:P13)" in body
    # (e) every error value
    for code, _m, _f in excel_io.SKILL_ERRORS:
        assert code in body, code
    # (f) named cells behave like absolute references
    assert "=F2*LABOR" in body and "never drifts" in body.lower()
    assert "RATE" in body and "HOLD" in body


def test_skills_sheet_is_readable():
    ws = _wb()[SHEET_SKILLS]
    for col in "ABCDEFGH":
        assert ws.column_dimensions[col].width >= 10
    assert ws["A1"].font.bold and ws["A1"].font.size > 11
    head, _first, _last = _exercise_rows(ws)
    assert ws[f"D{head}"].font.bold                  # table headings are bold
    assert all(len(str(ws[f"A{r}"].value or "")) < 200 for r in range(1, ws.max_row + 1))


# --------------------------------------------------------------------------- the exercise
def test_exercise_quantities_rate_and_prefilled_correct_answers():
    ws = _wb()[SHEET_SKILLS]
    head, first, last = _exercise_rows(ws)
    assert ws[f"J{head}"].value == excel_io.SKILLS_RATE
    for i, q in enumerate(excel_io.SKILLS_QUANTITIES):
        r = first + i
        assert ws[f"E{r}"].value == q
        assert ws[f"G{r}"].value == pytest.approx(q * excel_io.SKILLS_RATE)
    assert last - first + 1 == len(excel_io.SKILLS_QUANTITIES) == 6
    assert ws[f"E{last + 1}"].value is None          # the table really stops there


def test_exercise_check_formula_compares_student_cell_to_correct_answer():
    ws = _wb()[SHEET_SKILLS]
    head, first, last = _exercise_rows(ws)
    for r in range(first, last + 1):
        f = ws[f"H{r}"].value
        assert ws[f"H{r}"].data_type == "f", "the Check column must be a live Excel formula"
        assert f.startswith("=IF(")
        assert f"ABS(F{r}-G{r})" in f                # student cell vs the correct answer
        assert str(excel_io.SKILLS_TOLERANCE) in f
        assert '"OK"' in f and "Not yet" in f


def test_exercise_input_cells_are_unlocked_and_empty_everything_else_locked():
    ws = _wb()[SHEET_SKILLS]
    head, first, last = _exercise_rows(ws)
    assert ws.protection.sheet and not ws.protection.password
    for r in range(first, last + 1):
        assert ws[f"F{r}"].value is None, "the student's column must start empty"
        assert ws[f"F{r}"].protection.locked is False
        for col in "DEGHIJ":
            assert ws[f"{col}{r}"].protection.locked is not False, f"{col}{r} should stay locked"
    assert ws[f"J{head}"].protection.locked is not False
    assert ws["A1"].protection.locked is not False


def test_exercise_rate_cell_has_an_empty_drift_path_so_the_dollar_lesson_shows():
    """Copying a relative formula down must slide off the rate cell onto EMPTY cells."""
    ws = _wb()[SHEET_SKILLS]
    head, first, last = _exercise_rows(ws)
    for offset in range(1, last - first + 1):
        assert ws[f"J{head + offset}"].value is None
    body = _text(ws)
    assert "collapse to 0" in body and "$ signs" in body


# --------------------------------------------------------------------------- import is unaffected
def test_workbook_with_skills_sheet_round_trips_clean():
    def fill(wb):
        for kind, sheet in (("chase", "Chase Plan"), ("level", "Level Plan")):
            ws = wb[sheet]
            from aggplan import engine
            truth = worksheet.truth_worksheet(engine.reference_plan(S, kind, whole=True)[1])
            for i in range(12):
                for c in worksheet.EDIT_COLS:
                    ws[f"{worksheet.col_letter(c)}{worksheet.ROW_BASE + i}"] = float(
                        truth[c].iloc[i])

    wb = _wb()
    fill(wb)
    buf = io.BytesIO()
    wb.save(buf)
    res = import_workbook(buf.getvalue(), S, kinds=("chase", "level"))
    assert res.ok, [str(m) for m in res.errors]
    assert not res.errors and not res.warnings
    assert not any(m.sheet == SHEET_SKILLS for m in res.messages)
    assert SHEET_SKILLS not in res.sheets


def test_student_typing_on_the_skills_sheet_never_affects_the_import():
    def mut(wb):
        ws = wb[SHEET_SKILLS]
        head, first, last = _exercise_rows(ws)
        for r in range(first, last + 1):
            ws[f"F{r}"] = f"=E{r}*J{head}"          # the WRONG (relative) answer
        ws["A200"] = "my own scratch notes"
    wb = _wb()
    mut(wb)
    buf = io.BytesIO()
    wb.save(buf)
    res = import_workbook(buf.getvalue(), S, kinds=("chase",))
    assert not res.errors
    assert not any(m.sheet == SHEET_SKILLS for m in res.messages)


def test_legend_and_chart_live_right_of_column_q_so_the_importer_ignores_them():
    wb = _wb()
    q = worksheet.WS_COLS.index("Total Monthly Cost") + 1              # 17 = column Q
    for name in PLANS:
        ws = wb[name]
        assert ws["S1"].value and ws["R1"].value is None               # clearly separated
        # the importer reads A..Q only, and rejects anything it finds in rows 15..60 there
        for r in range(15, excel_io.MAX_ROWS + 1):
            for c in range(1, q + 1):
                assert ws.cell(row=r, column=c).value is None
        # the chart is anchored below the legend, well clear of the grid
        anchor = ws._charts[0].anchor._from
        assert anchor.col >= q and anchor.row >= 7
    res = import_workbook(build_template(S), S, kinds=("chase", "level"))
    assert not res.errors


# --------------------------------------------------------------------------- conditional formatting
def _rules(ws):
    return {str(rng.sqref): list(rng.rules) for rng in ws.conditional_formatting}


def test_conditional_formatting_rules_on_both_plan_sheets():
    wb = _wb()
    maxinv = float(S.params["max_inventory"])
    for name in PLANS:
        rules = _rules(wb[name])
        assert set(rules) == {"J2:J13", "K2:K13", "Q2:Q13"}

        back = rules["K2:K13"][0]                   # Backorders > 0 -> red
        assert back.type == "cellIs" and back.operator == "greaterThan"
        assert back.formula == ["0"]
        assert back.dxf.fill.bgColor.rgb.endswith("FFC7CE")

        inv = rules["J2:J13"][0]                    # Ending Inventory > max_inventory -> amber
        assert inv.type == "cellIs" and inv.operator == "greaterThan"
        assert float(inv.formula[0]) == pytest.approx(maxinv)
        assert inv.dxf.fill.bgColor.rgb.endswith("FFEB9C")

        cost = rules["Q2:Q13"][0]                   # Total Monthly Cost -> data bar
        assert cost.type == "dataBar" and cost.dataBar is not None


def test_colour_legend_is_worded_not_just_coloured():
    wb = _wb()
    maxinv = float(S.params["max_inventory"])
    for name in PLANS:
        ws = wb[name]
        legend = "\n".join(str(ws.cell(row=r, column=c).value)
                           for r in range(1, 9) for c in (19, 20)
                           if ws.cell(row=r, column=c).value)
        assert "COLOUR KEY" in legend
        assert "Red fill in column K (Backorders)" in legend
        assert "Amber fill in column J (Ending Inventory)" in legend
        assert "Blue bar in column Q (Total Monthly Cost)" in legend
        assert f"{maxinv:,.0f}" in legend
        assert "Regular Production" in legend       # the chart caption
        assert "empty until then" in legend


# --------------------------------------------------------------------------- the chart
def test_line_chart_series_titles_and_axes():
    wb = _wb()
    for name in PLANS:
        ws = wb[name]
        assert len(ws._charts) == 1
        ch = ws._charts[0]
        assert ch.tagname == "lineChart"
        assert [s.val.numRef.f for s in ch.series] == [f"'{name}'!$B$2:$B$13",
                                                       f"'{name}'!$E$2:$E$13"]
        assert [s.tx.strRef.f for s in ch.series] == [f"'{name}'!B1", f"'{name}'!E1"]
        assert ch.series[0].cat.numRef.f == f"'{name}'!$A$2:$A$13"
        title = "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []) if r.t)
        assert "forecast demand vs regular production" in title.lower()
        assert name in title
        assert ch.x_axis.title is not None and ch.y_axis.title is not None


def test_chart_parts_are_written_into_the_xlsx_itself():
    data = build_template(S)
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert sum(n.startswith("xl/charts/chart") for n in names) == 2
    body = zipfile.ZipFile(io.BytesIO(data)).read("xl/charts/chart1.xml").decode()
    assert "$B$2:$B$13" in body and "$E$2:$E$13" in body and "lineChart" in body


# --------------------------------------------------------------------------- whole-file sanity
def test_workbook_reloads_after_a_save_cycle_and_stays_small():
    data = build_template(S, student_label="Pat")
    assert len(data) < 200 * 1024, f"{len(data) / 1024:.0f} KB"
    wb = load_workbook(io.BytesIO(data))
    buf = io.BytesIO()
    wb.save(buf)
    wb2 = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb2.sheetnames == wb.sheetnames
    assert wb2.sheetnames[1] == SHEET_SKILLS
    assert len(_rules(wb2["Chase Plan"])) == 3
    assert import_workbook(buf.getvalue(), S).ok


def test_read_me_points_at_the_skills_sheet():
    readme = _text(_wb()["Read Me"], max_col=1)
    assert SHEET_SKILLS in readme
    assert "never imported" in readme and "never graded" in readme


def test_skills_sheet_is_present_in_every_template_variant():
    wb = _wb(hybrid=False, whole=False)
    assert wb.sheetnames[1] == SHEET_SKILLS
    assert "Hybrid Plan" not in wb.sheetnames
    res = import_workbook(build_template(S, hybrid=False, whole=False), S, whole=False)
    assert res.ok, [str(m) for m in res.errors]
