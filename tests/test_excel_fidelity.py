"""Two things the evaluation of this app found wrong about its Excel story:

1. the checker rejected a range that runs ACROSS a row, so `=SUM(L2:P2)` - the ordinary way
   to total a row of costs in Excel - was marked wrong, even though the workbook's own Column
   Guide describes that column as "Labor + Hiring + Layoff + Holding + Backorder";
2. nothing recorded WHAT the student typed, so a plan hand-calculated and entered as 180
   literal numbers was indistinguishable from a linked spreadsheet model, in the app and in
   the PDF alike.
"""

import io

import openpyxl
import pytest

from aggplan import excel_io
from aggplan.engine import reference_plan
from aggplan.formula import MAX_RANGE_CELLS, column_index, evaluate_grid
from aggplan.report_pdf import build_pdf
from aggplan.scenario import resolve_scenario
from aggplan.worksheet import (EDIT_COLS, ROW_BASE, WS_COLS, authoring_stats, blank_worksheet,
                               entry_rows, grade_worksheet, truth_worksheet)
from tests.fixtures import make_snapshot
from tests.helpers import SEASONAL, scn

COSTS = ["Regular Labor Cost", "Hiring Cost", "Layoff Cost", "Holding Cost", "Backorder Cost"]


def _grid(formula, extra=None):
    s = scn(SEASONAL)
    df = blank_worksheet(s)
    for i in (0, 1):
        for c, v in zip(COSTS, ["16000", "600", "0", "25", "0"]):
            df.at[ROW_BASE + i, c] = v
    for k, v in (extra or {}).items():
        df.at[ROW_BASE, k] = v
    df.at[ROW_BASE, "Total Monthly Cost"] = formula
    r = evaluate_grid(df, s.named_constants(), row_base=ROW_BASE)
    return r.values["Total Monthly Cost"][0], r.errors.get(("Total Monthly Cost", 0))


# ------------------------------------------------------------------ ranges
def test_sum_across_a_row_is_accepted():
    """The regression: in Excel this is THE way to total a row."""
    val, err = _grid("=SUM(L2:P2)")
    assert err is None and val == pytest.approx(16625.0)


def test_row_range_matches_writing_the_additions_out():
    a, _ = _grid("=SUM(L2:P2)")
    b, _ = _grid("=L2+M2+N2+O2+P2")
    assert a == b


@pytest.mark.parametrize("f,expected", [
    ("=SUM(L2:L3)", 32000.0),            # down a column (worked before)
    ("=SUM(L2:P2)", 16625.0),            # across a row
    ("=SUM(P2:L2)", 16625.0),            # written backwards
    ("=SUM(L2:P3)", 33250.0),            # a block
    ("=MAX(L2:P2)", 16000.0),
    ("=MIN(L2:P2)", 0.0),
    ("=SUM(L2:P2)/2", 8312.5),
    ("=ROUND(SUM(L2:P2),0)", 16625.0),
])
def test_range_shapes(f, expected):
    val, err = _grid(f)
    assert err is None, (f, err and err.message)
    assert val == pytest.approx(expected)


def test_ranges_still_report_real_problems():
    _, err = _grid("=SUM(A2:P2)")                   # Month column is text
    assert err is not None and err.code == "TYPE"
    for f in ("=SUM(L2:ZZ2)", "=SUM(L2:L99)"):      # outside the worksheet
        _, err = _grid(f)
        assert err is not None and err.code == "REF", (f, err and err.message)
        assert "isn't in the worksheet" in err.message
    _, err = _grid("=SUM(B2:Q13)")                  # in bounds but far too many cells
    assert err is not None and err.code == "LIMIT" and str(MAX_RANGE_CELLS) in err.message
    _, err = _grid("=SUM(Q2:Q2)")                   # a cell summing itself
    assert err is not None and err.code == "CIRC"


def test_column_index_is_the_inverse_of_column_letter():
    from aggplan.formula import column_letter
    for i in range(26):
        assert column_index(column_letter(i)) == i


def test_excel_workbook_built_with_formulas_is_accepted():
    """End to end: the failure that prompted this change."""
    s = scn(SEASONAL)
    dec, plan = reference_plan(s, "chase", policy_key="use_first", whole=True)
    truth = truth_worksheet(plan)
    F = {"Regular Labor Cost": "=F{r}*LABOR", "Hiring Cost": "=H{r}*HIRE",
         "Layoff Cost": "=I{r}*LAYOFF", "Holding Cost": "=J{r}*HOLD",
         "Backorder Cost": "=K{r}*BACKORDER", "Total Monthly Cost": "=SUM(L{r}:P{r})",
         "Workers": "=ROUNDUP(E{r}/RATE,0)", "Paid Labor Hours/Day": "=F{r}*SHIFT",
         "Ending Inventory": "=MAX(0,C{r}+E{r}-D{r}-B{r})",
         "Backorders": "=MAX(0,B{r}+D{r}-C{r}-E{r})"}
    wb = openpyxl.load_workbook(io.BytesIO(excel_io.build_template(s, whole=True, hybrid=True)))
    ws = wb["Chase Plan"]
    for i in range(12):
        r = i + 2
        for ci, col in enumerate(WS_COLS):
            if col in EDIT_COLS:
                ws.cell(r, ci + 1).value = (F[col].format(r=r) if col in F
                                            else float(truth[col].iloc[i]))
    buf = io.BytesIO()
    wb.save(buf)
    res = excel_io.import_workbook(buf.getvalue(), s, whole=True, kinds=("chase",))
    assert res.ok and not [m for m in res.messages if m.level == "error"]
    g = grade_worksheet(res.sheets["chase"], s, truth, kind="chase", whole=True)
    assert g.complete and g.n_wrong == 0
    # the formulas survive import as text, they are not flattened to numbers
    assert str(res.sheets["chase"].iloc[0]["Total Monthly Cost"]).startswith("=SUM(")


# ------------------------------------------------- what the student typed
def test_entry_rows_preserve_exactly_what_was_typed():
    s = scn(SEASONAL)
    ws = {c: [""] * 12 for c in EDIT_COLS}
    ws["Total Monthly Cost"][0] = "=SUM(L2:P2)"
    ws["Workers"][0] = " =ROUNDUP(E2/RATE) "
    ws["Hires"][0] = "3"
    rows = entry_rows(ws, list(s.months))
    assert len(rows) == 12 and rows[0]["Month"] == "January"
    assert rows[0]["Total Monthly Cost"] == "=SUM(L2:P2)"
    assert rows[0]["Workers"] == "=ROUNDUP(E2/RATE)"        # trimmed, not rewritten
    assert rows[0]["Hires"] == "3" and rows[0]["Layoffs"] == ""


def test_authoring_distinguishes_a_model_from_typed_numbers():
    s = scn(SEASONAL)
    literals = {c: [str(100 + i) for i in range(12)] for c in EDIT_COLS}
    a = authoring_stats(literals, s.named_constants())
    assert a["formula_cells"] == 0 and a["literal_cells"] == 180
    assert a["formula_share_pct"] == 0 and a["named_constants_used"] == []

    modelled = dict(literals)
    modelled["Regular Labor Cost"] = ["=F%d*LABOR" % (i + 2) for i in range(12)]
    modelled["Total Monthly Cost"] = ["=SUM(L%d:P%d)" % (i + 2, i + 2) for i in range(12)]
    b = authoring_stats(modelled, s.named_constants())
    assert b["formula_cells"] == 24 and b["literal_cells"] == 156
    assert b["cells_referencing_other_cells"] == 24
    assert "LABOR" in b["named_constants_used"]
    assert b["formula_share_pct"] > a["formula_share_pct"]


def test_authoring_counts_blanks_and_survives_missing_state():
    s = scn(SEASONAL)
    empty = authoring_stats(None, s.named_constants())
    assert empty["filled_cells"] == 0 and empty["formula_share_pct"] == 0
    partial = authoring_stats({"Hires": ["1", "", "2"]}, s.named_constants())
    assert partial["filled_cells"] == 2 and partial["blank_cells"] == 1


def test_snapshot_carries_the_entries_and_authoring():
    snap = make_snapshot()
    ch = snap["plans"]["chase"]
    assert "entries" in ch and "authoring" in ch and "entry_columns" in ch
    assert len(ch["entries"]) == 12
    assert set(ch["entry_columns"]) == set(EDIT_COLS)


def test_pdf_shows_the_method_and_the_exact_entries():
    from pypdf import PdfReader
    snap = make_snapshot()
    snap["plans"]["chase"]["entries"][0]["Total Monthly Cost"] = "=SUM(L2:P2)"
    snap["plans"]["chase"]["authoring"].update(formula_cells=24, literal_cells=156,
                                               filled_cells=180, formula_share_pct=13.3,
                                               named_constants_used=["LABOR", "RATE"])
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(build_pdf(snap))).pages)
    assert "How the worksheets were built" in text
    assert "Share built with formulas" in text and "13%" in text
    assert "entries exactly as the student typed" in text
    assert "=SUM(L2:P2)" in text.replace(" ", "")       # verbatim, not reformatted
    assert "Quantities" in text and "Costs" in text
    # the method is reported neutrally, never as a mark
    assert "describe the method, not its correctness" in text.replace("\n", " ")


def test_pdf_handles_a_plan_with_no_recorded_entries():
    snap = make_snapshot()
    snap["plans"]["level"]["entries"] = []
    assert build_pdf(snap).startswith(b"%PDF")


def test_entry_text_is_escaped_not_interpreted():
    from pypdf import PdfReader
    snap = make_snapshot()
    snap["plans"]["chase"]["entries"][0]["Hires"] = '<b>=A1&B1</b>'
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(build_pdf(snap))).pages)
    assert "<b>" in text
