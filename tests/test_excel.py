"""Excel template / import / export tests (aggplan.excel_io)."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import load_workbook

from aggplan import MODEL_VERSION, engine, excel_io, text, worksheet
from aggplan.excel_io import build_template, export_workbook, import_workbook
from tests.helpers import SEASONAL, scn

S = scn(SEASONAL)


# --------------------------------------------------------------------------- helpers
def _truth(kind, scenario=S, whole=True):
    dec, df = engine.reference_plan(scenario, kind, whole=whole)
    return worksheet.truth_worksheet(df)


def _fill(wb, sheet, truth, formulas=False):
    """Write the app's truth into C2:Q13; optionally use formulas for some columns."""
    ws = wb[sheet]
    L = worksheet.col_letter
    for i in range(12):
        r = worksheet.ROW_BASE + i
        for c in worksheet.EDIT_COLS:
            ws[f"{L(c)}{r}"] = float(truth[c].iloc[i])
        if formulas:
            ws[f"{L('Paid Labor Hours/Day')}{r}"] = f"={L('Workers')}{r}*SHIFT"
            ws[f"{L('Regular Labor Cost')}{r}"] = f"={L('Workers')}{r}*LABOR"
            ws[f"{L('Holding Cost')}{r}"] = f"={L('Ending Inventory')}{r}*HOLD"
            ws[f"{L('Workers')}{r}"] = f"=ROUNDUP({L('Regular Production')}{r}/RATE,0)" \
                if sheet == "Chase Plan" else float(truth["Workers"].iloc[i])


def _roundtrip(mutator=None, **kw):
    wb = load_workbook(io.BytesIO(build_template(S, **kw)))
    if mutator:
        mutator(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _grade(res, kind, scenario=S):
    return worksheet.grade_worksheet(res.sheets[kind], scenario, _truth(kind, scenario),
                                     kind=kind, whole=True)


def _msgs(res, level):
    return [m for m in res.messages if m.level == level]


# --------------------------------------------------------------------------- template
def test_template_structure_and_identification():
    wb = load_workbook(io.BytesIO(build_template(S, student_label="Pat")))
    assert wb.sheetnames[0] == "Read Me"
    for name in ("Scenario", "Chase Plan", "Level Plan", "Hybrid Plan", "Compare", "_meta"):
        assert name in wb.sheetnames
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(min_row=2, max_col=2)}
    assert meta["schema"] == 1 and meta["model_version"] == MODEL_VERSION
    assert meta["seed"] == S.seed and meta["config_hash"] == S.config_hash
    assert meta["worker_model"] == "whole" and meta["variant"] == 0
    readme = "\n".join(str(r[0].value) for r in wb["Read Me"].iter_rows() if r[0].value)
    assert S.config_hash in readme and MODEL_VERSION in readme and str(S.seed) in readme
    assert "Safety stock = a PLANNING TARGET" in readme          # engine rules copied
    assert "Terminal requirements" in readme


def test_plan_sheet_layout():
    wb = load_workbook(io.BytesIO(build_template(S)))
    for name in ("Chase Plan", "Level Plan"):
        ws = wb[name]
        assert [ws.cell(row=1, column=j + 1).value for j in range(17)] == worksheet.WS_COLS
        for i in range(12):
            assert ws.cell(row=2 + i, column=1).value == S.months[i]
            assert ws.cell(row=2 + i, column=2).value == S.demand[i]
            assert all(ws.cell(row=2 + i, column=j).value is None for j in range(3, 18))
        assert ws.freeze_panes == "C2"
        assert ws.protection.sheet and not ws.protection.password
        assert ws["C2"].protection.locked is False and ws["Q13"].protection.locked is False
        assert ws["A2"].protection.locked and ws["B2"].protection.locked
        assert ws["A14"].value.startswith("Σ")
        for c in worksheet.SUM_COLS:
            L = worksheet.col_letter(c)
            assert ws[f"{L}14"].value == f"=SUM({L}2:{L}13)"
        assert ws["E1"].comment is not None and "bottles" in ws["E1"].comment.text
    hy = wb["Hybrid Plan"]
    assert [hy.cell(row=1, column=j + 1).value for j in range(6)] == excel_io.HYBRID_IN_COLS
    assert hy["W1"].value and any("Subcontract capacity" == hy[f"W{r}"].value for r in range(2, 12))


def test_template_without_hybrid_and_partial_units():
    wb = load_workbook(io.BytesIO(build_template(S, hybrid=False, whole=False)))
    assert "Hybrid Plan" not in wb.sheetnames
    guide = "\n".join(str(c.value) for row in wb["Column Guide"].iter_rows() for c in row if c.value)
    assert "worker-equivalents" in guide and "hours/day of capacity added" in guide
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(min_row=2, max_col=2)}
    assert meta["worker_model"] == "partial"
    assert "Hybrid Plan" not in str(meta["plan_sheets"])
    assert "worker-equivalents" in wb["Chase Plan"]["F1"].comment.text


def test_defined_names_resolve_to_constants():
    wb = load_workbook(io.BytesIO(build_template(S)))
    consts = S.named_constants()
    assert set(consts) <= set(wb.defined_names.keys())
    for name, val in consts.items():
        dest = list(wb.defined_names[name].destinations)
        assert len(dest) == 1
        sheet, ref = dest[0]
        assert sheet == "Scenario"
        assert wb[sheet][ref.replace("$", "")].value == pytest.approx(val)
        assert excel_io._valid_defined_name(name)
    assert not excel_io._valid_defined_name("A1") and not excel_io._valid_defined_name("R1C1")


def test_scenario_sheet_numbers_are_numeric_and_protected():
    ws = load_workbook(io.BytesIO(build_template(S)))["Scenario"]
    assert ws.protection.sheet and not ws.protection.password
    assert [ws[f"B{6 + i}"].value for i in range(12)] == list(S.demand)
    labels = {r[0].value: r[1].value for r in ws.iter_rows(min_row=20, max_col=2)
              if r[0].value}
    assert labels["Capacity per worker per month"] == S.cap_month
    assert labels["Holding"] == 0.25
    assert labels["Overtime limit"] == 0.2
    assert ws["B6"].font.color.rgb.endswith("0000FF")


# --------------------------------------------------------------------------- import
def test_roundtrip_numbers_and_formulas_grade_ok():
    def fill(wb):
        _fill(wb, "Chase Plan", _truth("chase"), formulas=True)
        _fill(wb, "Level Plan", _truth("level"))
    res = import_workbook(_roundtrip(fill), S, kinds=("chase", "level"))
    assert res.ok, [str(m) for m in res.errors]
    for kind in ("chase", "level"):
        df = res.sheets[kind]
        assert list(df.columns) == worksheet.WS_COLS
        assert list(df.index) == list(range(2, 14))
        g = _grade(res, kind)
        assert g.n_wrong == 0 and g.n_blank == 0 and g.n_ok == 12 * 15
    chase = res.sheets["chase"]
    assert chase.at[2, "Paid Labor Hours/Day"] == "=F2*SHIFT"      # formulas survive as text
    assert chase.at[5, "Workers"].startswith("=ROUNDUP(")
    assert chase.at[2, "Month"] == "January" and chase.at[2, "Forecast Demand"] == S.demand[0]
    assert res.hybrid is None
    assert not _msgs(res, "warning")


def test_blank_sheets_warn_only_for_requested_kinds():
    res = import_workbook(build_template(S), S, kinds=("chase",))
    assert res.ok and "chase" in res.sheets and "level" not in res.sheets
    assert any("completely blank" in m.text for m in _msgs(res, "warning"))
    assert not any(m.sheet == "Level Plan" for m in res.messages)


def test_partial_blank_counts():
    def fill(wb):
        _fill(wb, "Chase Plan", _truth("chase"))
        wb["Chase Plan"]["C2"] = None
        wb["Chase Plan"]["D3"] = None
    res = import_workbook(_roundtrip(fill), S, kinds=("chase",))
    assert res.ok
    assert any("2 of 180" in m.text for m in _msgs(res, "warning"))


def test_scenario_mismatch_is_blocking_with_useful_message():
    other = scn(SEASONAL, seed=99, holding_cost=0.5)
    res = import_workbook(build_template(S), other)
    assert not res.ok
    msg = res.errors[0].text
    assert f"seed {S.seed} (config {S.config_hash})" in msg
    assert f"seed {other.seed} (config {other.config_hash})" in msg
    assert "Download a new template for your current scenario" in msg


def test_worker_model_mismatch_blocked():
    res = import_workbook(build_template(S, whole=False), S, whole=True)
    assert not res.ok and "partial" in res.errors[0].text


def test_tampered_heading_names_cell_and_expected():
    def mut(wb):
        wb["Level Plan"]["F1"] = "Staff"
    res = import_workbook(_roundtrip(mut), S, kinds=("level",))
    assert not res.ok and "level" not in res.sheets
    m = res.errors[0]
    assert m.cell == "F1" and "'Workers'" in m.text and "'Staff'" in m.text


def test_tampered_demand_and_month():
    def mut(wb):
        wb["Chase Plan"]["B5"] = 1234
        wb["Chase Plan"]["A3"] = "Febuary"
    res = import_workbook(_roundtrip(mut), S, kinds=("chase",))
    cells = {m.cell: m.text for m in res.errors}
    assert not res.ok
    assert "B5" in cells and f"{S.demand[3]:,.0f}" in cells["B5"] and "1,234" in cells["B5"]
    assert "A3" in cells and "'February'" in cells["A3"]


def test_text_in_numeric_cell_and_outside_cells():
    def mut(wb):
        wb["Chase Plan"]["E4"] = "lots"
        wb["Chase Plan"]["C20"] = 5
        wb["Chase Plan"]["R2"] = "ignored note"          # right of Q: ignored
        wb["Chase Plan"]["E14"] = 999                    # totals row: ignored
    res = import_workbook(_roundtrip(mut), S, kinds=("chase",))
    errs = {m.cell: m.text for m in res.errors}
    assert not res.ok
    assert "E4" in errs and "'lots'" in errs["E4"]
    assert "C20" in errs and "outside" in errs["C20"]
    assert set(errs) == {"E4", "C20"}


def test_missing_sheet_and_meta():
    def drop_level(wb):
        del wb["Level Plan"]
    res = import_workbook(_roundtrip(drop_level), S, kinds=("level",))
    assert not res.ok and "Level Plan" in res.errors[0].text

    def drop_meta(wb):
        del wb["_meta"]
    res = import_workbook(_roundtrip(drop_meta), S)
    assert not res.ok and "_meta" in res.errors[0].text

    def bad_schema(wb):
        wb["_meta"]["B2"] = 7
    assert not import_workbook(_roundtrip(bad_schema), S).ok

    def bad_version(wb):
        for r in wb["_meta"].iter_rows(min_row=2, max_col=2):
            if r[0].value == "model_version":
                r[1].value = "2.9.0"
    res = import_workbook(_roundtrip(bad_version), S)
    assert not res.ok and "model version" in res.errors[0].text


def test_hybrid_missing_from_template_is_info_not_error():
    res = import_workbook(build_template(S, hybrid=False), S)
    assert res.ok and res.hybrid is None
    assert any(m.level == "info" and "Hybrid" in m.sheet for m in res.messages)


def test_hybrid_import():
    def mut(wb):
        ws = wb["Hybrid Plan"]
        for i in range(12):
            r = 2 + i
            ws[f"C{r}"] = 10
            ws[f"D{r}"] = f"=C{r}*RATE"
            if i == 5:
                ws[f"E{r}"] = 1500
            if i == 6:
                ws[f"F{r}"] = 2000
        ws["I3"] = "my note"                              # calculation columns are lenient
    res = import_workbook(_roundtrip(mut), S, kinds=("hybrid",))
    assert res.ok, [str(m) for m in res.errors]
    h = res.hybrid
    assert h["workers"] == [10.0] * 12 and h["regular"] == [10000.0] * 12
    assert h["overtime"][5] == 1500.0 and h["overtime"][0] is None
    assert h["subcontract"][6] == 2000.0


def test_hybrid_text_rejected():
    def mut(wb):
        wb["Hybrid Plan"]["D2"] = "abc"
    res = import_workbook(_roundtrip(mut), S, kinds=("hybrid",))
    assert not res.ok and res.errors[0].cell == "D2" and res.hybrid is None


def test_unsupported_formula_without_cache_warns_and_keeps_formula():
    def mut(wb):
        _fill(wb, "Chase Plan", _truth("chase"))
        wb["Chase Plan"]["K5"] = "=IF(1,2,3)"            # openpyxl writes no cached value
    res = import_workbook(_roundtrip(mut), S, kinds=("chase",))
    assert res.ok
    w = [m for m in _msgs(res, "warning") if m.cell == "K5"]
    assert w and "IF" in w[0].text and "MAX/MIN" in w[0].text
    assert res.sheets["chase"].at[5, "Backorders"] == "=IF(1,2,3)"


def _xlsxwriter_book(cell_formula, cached):
    """A workbook with a REAL cached value (xlsxwriter can write one)."""
    import xlsxwriter
    bio = io.BytesIO()
    wb = xlsxwriter.Workbook(bio, {"in_memory": True})
    meta = wb.add_worksheet("_meta")
    rows = [("schema", 1), ("model_version", MODEL_VERSION), ("seed", S.seed),
            ("variant", S.variant), ("config_hash", S.config_hash), ("worker_model", "whole"),
            ("plan_sheets", "Chase Plan")]
    for r, (k, v) in enumerate(rows, start=1):
        meta.write(r, 0, k)
        meta.write(r, 1, v)
    ws = wb.add_worksheet("Chase Plan")
    for j, name in enumerate(worksheet.WS_COLS):
        ws.write(0, j, name)
    for i in range(12):
        ws.write(1 + i, 0, S.months[i])
        ws.write(1 + i, 1, S.demand[i])
    ws.write_number(1, 4, 10000)                         # E2
    ws.write_formula(cell_formula[0], cell_formula[1], None, cached)   # F2 + cached value
    wb.close()
    return bio.getvalue()


def test_unsupported_formula_with_cached_value_falls_back():
    data = _xlsxwriter_book(("F2", "=IF(E2>0,ROUNDUP(E2/RATE,0),0)"), 10)
    res = import_workbook(data, S, kinds=("chase",))
    assert res.ok, [str(m) for m in res.errors]
    assert res.sheets["chase"].at[2, "Workers"] == "10"
    info = [m for m in _msgs(res, "info") if m.cell == "F2"]
    assert info and "Excel's calculated value" in info[0].text


def test_excel_error_value_is_error():
    data = _xlsxwriter_book(("F2", "=1/0"), "#DIV/0!")
    res = import_workbook(data, S, kinds=("chase",))
    assert not res.ok
    m = [m for m in res.errors if m.cell == "F2"]
    assert m and "#DIV/0!" in m[0].text


def test_corrupt_and_oversize_inputs_never_raise():
    for junk in (b"", b"not a zip at all", b"PK\x03\x04garbage" * 10, None):
        res = import_workbook(junk, S)
        assert res.ok is False and res.errors
    good = build_template(S)
    res = import_workbook(good[: len(good) // 2], S)
    assert res.ok is False and res.errors
    big = b"\0" * (6 * 1024 * 1024)
    res = import_workbook(big, S)
    assert not res.ok and "MB" in res.errors[0].text


def test_zip_bomb_rejected():
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/big.xml", b"0" * (60 * 1024 * 1024))
    data = bio.getvalue()
    assert len(data) < excel_io.MAX_BYTES
    res = import_workbook(data, S)
    assert not res.ok and "50 MB" in res.errors[0].text


# --------------------------------------------------------------------------- export
def _plan_bundle(kind, whole=True):
    dec, df = engine.reference_plan(S, kind, whole=whole)
    req = engine.requirements_for(S, dec.policy)
    return {"df": df, "summary": engine.summarize(df, S, dec.policy),
            "checks": engine.check_feasibility(df, req),
            "settings": {"Worker model": "Whole workers", "Inventory policy": "=use first"}}


def test_export_totals_and_summary_match_engine():
    plans = {"Chase": _plan_bundle("chase"), "Level": _plan_bundle("level")}
    data = export_workbook(S, plans, student_label="Pat", status="Draft")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[:2] == ["Summary", "Assumptions"]
    for lab, b in plans.items():
        ws = wb[f"{lab} Plan"]
        assert [c.value for c in ws[1]] == engine.PLAN_COLUMNS
        df = b["df"]
        for j, col in enumerate(engine.PLAN_COLUMNS, start=1):
            if col in ("Month", "Storage Heuristic"):
                continue
            for i in range(12):
                assert ws.cell(row=2 + i, column=j).value == float(df[col].iloc[i])   # exact
        tot = ws[14]
        for col in excel_io._SUM_EXPORT:
            j = engine.PLAN_COLUMNS.index(col) + 1
            L = ws.cell(row=1, column=j).column_letter
            assert tot[j - 1].value == f"=SUM({L}2:{L}13)"
            stored = sum(ws.cell(row=2 + i, column=j).value for i in range(12))
            assert stored == pytest.approx(float(df[col].sum()), abs=1e-6)
        total_col = engine.PLAN_COLUMNS.index("Total Monthly Cost") + 1
        assert sum(ws.cell(row=2 + i, column=total_col).value for i in range(12)) == \
            pytest.approx(b["summary"]["Total cost"], abs=0.005)
        assert ws.freeze_panes == "B2"
    # Summary sheet values + definitions + plain-text feasibility
    sm = wb["Summary"]
    rows = {r[0].value: [c.value for c in r] for r in sm.iter_rows() if r[0].value}
    assert rows["Total cost"][1] == plans["Chase"]["summary"]["Total cost"]
    assert rows["Total cost"][2] == plans["Level"]["summary"]["Total cost"]
    assert rows["On-time fulfillment"][3] == engine.METRIC_DEFINITIONS["On-time fulfillment"]
    feas = rows["Feasibility"]
    for k, lab in enumerate(("Chase", "Level"), start=1):
        ok = engine.is_feasible(plans[lab]["checks"])
        assert (feas[k] == "Meets all requirements") if ok else feas[k].startswith("Does NOT meet:")
    alltext = "\n".join(str(c.value) for row in wb["Assumptions"].iter_rows() for c in row
                        if c.value is not None)
    assert S.config_hash in alltext and MODEL_VERSION in alltext
    # a setting that starts with '=' is stored as text, not a formula
    cell = [c for row in sm.iter_rows() for c in row if c.value == "=use first"]
    assert cell and cell[0].data_type == "s"


def test_export_failing_plan_is_described_in_text():
    b = _plan_bundle("level")
    b["checks"] = list(b["checks"]) + [{"key": "x", "label": "Storage limit", "ok": False,
                                        "detail": "Storage limit exceeded in May"}]
    wb = load_workbook(io.BytesIO(export_workbook(S, {"Level": b})))
    feas = [r for r in wb["Summary"].iter_rows(values_only=True) if r[0] == "Feasibility"][0]
    assert feas[1].startswith("Does NOT meet:") and "Storage limit" in feas[1]


@pytest.mark.skipif(__import__("shutil").which("soffice") is None, reason="LibreOffice not installed")
def test_libreoffice_recalculates_template(tmp_path):      # optional extra, never required
    import subprocess
    p = tmp_path / "t.xlsx"
    p.write_bytes(_roundtrip(lambda wb: _fill(wb, "Chase Plan", _truth("chase"), formulas=True)))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
                    str(tmp_path / "out"), str(p)], check=True, timeout=120)
    res = import_workbook((tmp_path / "out" / "t.xlsx").read_bytes(), S, kinds=("chase",))
    assert res.ok
